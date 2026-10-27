"""Native editable Office Math, with centered expressions and end numbers."""
import re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.enum.text import WD_TAB_ALIGNMENT


def elem(tag,*children):
    node=OxmlElement('m:'+tag)
    for child in children:node.append(child)
    return node


def run(text,upright=False):
    node=elem('r')
    if upright:
        style=elem('sty');style.set(qn('m:val'),'p');node.append(elem('rPr',style))
    value=elem('t');value.text=text;node.append(value);return node


def sub(value,index):return elem('sSub',elem('e',value),elem('sub',run(index)))
def sup(value,index):return elem('sSup',elem('e',value),elem('sup',run(index)))
def frac(top,bottom):return elem('f',elem('num',*top),elem('den',*bottom))


def hat(value):
    mark=elem('chr');mark.set(qn('m:val'),'̂')
    return elem('acc',elem('accPr',mark),elem('e',value))


def summation(lower,upper,content):
    char=elem('chr');char.set(qn('m:val'),'∑');location=elem('limLoc');location.set(qn('m:val'),'undOvr')
    return elem('nary',elem('naryPr',char,location),elem('sub',run(lower)),elem('sup',run(upper)),elem('e',*content))


def delimiter(content,left='(',right=')'):
    before=elem('begChr');before.set(qn('m:val'),left);after=elem('endChr');after.set(qn('m:val'),right)
    return elem('d',elem('dPr',before,after),elem('e',*content))


def formula(key):
    date=lambda index:sub(hat(run('d')),index)
    if key=='process':
        positive=sub(delimiter([sub(run('L'),'τ'),run(' − 12.5')],'[',']'),'+')
        term=frac([run('r'),delimiter([sub(run('T'),'τ')])],
                  [run('1 + b '),positive,run(' '),sub(run('I'),'τ')])
        return [run('P(t) = '),summation('τ=1','t',[term])]
    if key=='transformer':
        return [run('D(t) = '),summation('τ=1','t',[run('exp',True),delimiter([run('u'),delimiter([sub(run('h'),'τ')])])]),
                run(',   '),date('j'),run(' = '),run('min',True),delimiter([run('t : D(t) ≥ '),sub(run('κ'),'j')],'{','}')]
    if key=='residual':
        return [hat(run('d')),run(' = Π'),delimiter([sup(hat(run('d')),'T–P'),run(' + λ g(H, c)')])]
    if key=='silking':
        return [date('S'),run(' = '),date('T'),run(' + α '),delimiter([date('M'),run(' − '),date('T')])]
    if key=='mae':
        absolute=delimiter([date('ijs'),run(' − '),sub(run('d'),'ij')],'|','|')
        inner=summation('i=1','Nⱼ',[absolute])
        stage=frac([run('1')],[sub(run('N'),'j')])
        return [run('MAE',True),run(' = '),frac([run('1')],[run('3J')]),run(' '),
                summation('s=1','3',[summation('j=1','J',[stage,run(' '),inner])])]
    if key=='cardinal':
        response=sub(delimiter([run('2u(T) − '),sup(run('u(T)'),'2')],'[',']'),'+')
        return [run('r(T) = '),delimiter([sub(run('T'),'opt'),run(' − 8')]),run(' '),response]
    if key=='cardinal_shape':
        u=sup(delimiter([frac([run('T − 8')],[sub(run('T'),'opt'),run(' − 8')])]),'a')
        a=frac([run('ln 2',True)],[run('ln',True),delimiter([frac([sub(run('T'),'max'),run(' − 8')],[sub(run('T'),'opt'),run(' − 8')])])])
        return [run('u(T) = '),u,run(',   a = '),a]
    raise ValueError('Unknown professional equation: '+key)


def add_equation(doc,description):
    match=re.fullmatch(r'([a-z_]+) \((\d+)\)',description.strip())
    if match is None:raise ValueError('Equation key and end number required')
    key,number=match.groups();math=elem('oMath',*formula(key))
    paragraph=doc.add_paragraph();paragraph.paragraph_format.keep_together=True
    section=doc.sections[-1];width=section.page_width-section.left_margin-section.right_margin
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(width/914400/2),WD_TAB_ALIGNMENT.CENTER)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(width/914400),WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run('\t');paragraph._p.append(math);paragraph.add_run('\t('+number+')')
    return paragraph
